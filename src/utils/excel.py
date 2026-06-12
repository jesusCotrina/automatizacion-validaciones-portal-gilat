
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import date
from openpyxl.utils import get_column_letter

def generar_excel_formato(df,periodo):
    periodo_carpeta = periodo[:7]
    path_resumen= f"../resultados/{periodo_carpeta}/resumen_validaciones.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"

    # Fila 1
    ws["A1"] = "Fecha de actualización:"
    ws["B1"] = date.today().strftime("%Y-%m-%d")

    # Filas 2 y 3 vacías
    ws.append([])
    ws.append([])

    # El DataFrame comenzará en la fila 4
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)

    # =====================
    # CABECERAS
    # =====================

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )

    for cell in ws[4]:
        cell.fill = header_fill
        cell.font = header_font

    # =====================
    # TABLA EXCEL
    # =====================

    ultima_fila = ws.max_row
    ultima_columna = ws.max_column

    rango = f"A4:{get_column_letter(ultima_columna)}{ultima_fila}"

    table = Table(
        displayName="ResumenReportes",
        ref=rango
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    table.tableStyleInfo = style
    ws.add_table(table)

    # =====================
    # COLORES ESTADOS
    # =====================

    verde = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE"
    )

    amarillo = PatternFill(
        fill_type="solid",
        fgColor="FFEB9C"
    )

    rojo = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE"
    )

    for row in ws.iter_rows(min_row=2):

        for cell in row:

            valor = str(cell.value).upper() if cell.value else ""

            if valor == "VALIDADO":
                cell.fill = verde

            elif valor == "OBSERVADO":
                cell.fill = amarillo

            elif valor == "ERROR":
                cell.fill = rojo

    # =====================
    # AUTOAJUSTE COLUMNAS
    # =====================

    for column in ws.columns:

        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:

            try:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )
            except:
                pass

        ws.column_dimensions[column_letter].width = min(
            max_length + 3,
            60
        )

    wb.save(path_resumen)
    return path_resumen